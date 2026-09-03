#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
scrape_voz.py — Thu thap & xep hang thao luan trong box "Diem bao" cua voz.vn
================================================================================
Muc tieu (theo yeu cau):
  0. Chay hang ngay 10:00 (gio VN) -> xem .github/workflows/daily.yml
  1. Quet PAGE 1 cua https://voz.vn/f/diem-bao.33/  (BO QUA thread sticky/ghim mau cam)
  2. Vao chi tiet tung thread, quet HET cac trang cua thread do
  3. Loc top N comment co nhieu "Ung" nhat (mac dinh 20)
  4. Xuat Excel: link bai bao (lay tu post #1), tieu de, top comment 1..20
     + xep hang thread nao dang "hot" nhat

Ghi chu tuan thu (doc ky truoc khi dung):
  - robots.txt cua voz.vn cho "User-agent: *" la Allow: / (chi chan /admin.php)
    -> script nay duoc phep doc /f/ va /t/.
  - robots.txt CHAN cac bot AI co ten (GPTBot, ClaudeBot, CCBot, Bytespider...).
    Vi vay TUYET DOI khong dat User-Agent thanh ten cua cac bot do.
  - Content-Signal cua voz.vn: search=yes, ai-train=no, use=reference.
    -> Dung noi dung lam TU LIEU THAM KHAO / y tuong. Khong dung de train model,
       khong copy-paste nguyen van comment cua nguoi khac len Threads.
  - Rate limit lich su (mac dinh 1.2s/request) va User-Agent trung thuc.

Su dung:
    python scrape_voz.py                      # quet + xuat Excel vao ./output
    python scrape_voz.py --top 20 --delay 1.5
    python scrape_voz.py --max-threads 5      # test nhanh
    python scrape_voz.py --dry-run            # chi liet ke thread, khong vao chi tiet
"""
from __future__ import annotations

import argparse
import json
import logging
import os
import random
import re
import sys
import time
import urllib.parse as urlparse
import urllib.robotparser as robotparser
from dataclasses import dataclass, field, asdict
from datetime import datetime, timezone, timedelta
from typing import Iterable, Optional

import requests
from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

try:
    from zoneinfo import ZoneInfo
    TZ_VN = ZoneInfo("Asia/Bangkok")
except Exception:                                            # pragma: no cover
    TZ_VN = timezone(timedelta(hours=7))

# ------------------------------------------------------------------ constants
BASE = "https://voz.vn"
FORUM_PATH = "/f/diem-bao.33/"
FORUM_URL = BASE + FORUM_PATH

# UA thuong cua trinh duyet. KHONG dat ten bot AI (robots.txt cua voz chan ho).
DEFAULT_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/140.0.0.0 Safari/537.36"
)

# --- Cong thuc "hot score" (tunable) ---------------------------------------
# base  = replies * W_REPLY + tong_ung * W_UNG
# hot   = base / (tuoi_thread_gio + AGE_OFFSET_H) ** GRAVITY
# GRAVITY < 1 -> giam bot loi the cua thread moi tinh; AGE_OFFSET_H chong chia 0.
W_REPLY = 1.0
W_UNG = 2.0
GRAVITY = 0.8
AGE_OFFSET_H = 6.0

# Host khong phai "bai bao" -> bo qua khi do tim link bai bao trong post #1
NON_ARTICLE_HOSTS = {
    "voz.vn", "www.voz.vn", "next.voz.vn", "statics.voz.tech", "vozforum.org",
    "i.imgur.com", "imgur.com", "i.ibb.co", "ibb.co", "prnt.sc",
    "drive.google.com", "photos.google.com",
}
IMAGE_EXT_RE = re.compile(r"\.(png|jpe?g|gif|webp|bmp|svg)(\?|$)", re.I)
THREAD_URL_RE = re.compile(r"^/t/(?P<slug>[^/?#]*?\.(?P<id>\d+))(?:/|$)")

log = logging.getLogger("voz")


# ------------------------------------------------------------------ dataclasses
@dataclass
class Post:
    post_id: str
    number: Optional[int]
    author: str
    url: str
    created_at: Optional[datetime]
    reactions: int
    text: str

    def to_dict(self) -> dict:
        d = asdict(self)
        d["created_at"] = self.created_at.isoformat() if self.created_at else None
        return d


@dataclass
class Thread:
    thread_id: str
    title: str
    url: str
    prefix: str = ""
    op_author: str = ""
    op_created_at: Optional[datetime] = None
    article_url: str = ""
    article_source: str = ""
    article_excerpt: str = ""
    status: str = "ok"          # ok | gone | error
    pages_scanned: int = 0
    replies: int = 0
    total_reactions: int = 0
    base_score: float = 0.0
    hot_score: float = 0.0
    posts: list[Post] = field(default_factory=list)
    top_comments: list[Post] = field(default_factory=list)
    error: str = ""

    def to_dict(self) -> dict:
        return {
            "thread_id": self.thread_id,
            "title": self.title,
            "url": self.url,
            "prefix": self.prefix,
            "op_author": self.op_author,
            "op_created_at": self.op_created_at.isoformat() if self.op_created_at else None,
            "article_url": self.article_url,
            "article_source": self.article_source,
            "article_excerpt": self.article_excerpt,
            "status": self.status,
            "pages_scanned": self.pages_scanned,
            "replies": self.replies,
            "total_reactions": self.total_reactions,
            "base_score": round(self.base_score, 2),
            "hot_score": round(self.hot_score, 3),
            "top_comments": [p.to_dict() for p in self.top_comments],
            "error": self.error,
        }


# ------------------------------------------------------------------ fetch layer
class Fetcher:
    """HTTP GET co retry + rate limit lich su.

    Backend `requests` la mac dinh. Neu voz/Cloudflare tra 403 (thuong xay ra
    voi IP datacenter nhu GitHub Actions), cai `curl_cffi` va dung
    --impersonate chrome  -> gia lap TLS fingerprint cua Chrome.
    """

    def __init__(self, ua: str = DEFAULT_UA, delay: float = 1.2,
                 timeout: int = 25, impersonate: str = ""):
        self.delay = max(0.0, delay)
        self.timeout = timeout
        self.ua = ua
        self.impersonate = impersonate
        self._last = 0.0
        self.headers = {
            "User-Agent": ua,
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "vi,vi-VN;q=0.9,en-US;q=0.8,en;q=0.7",
            "Upgrade-Insecure-Requests": "1",
            "Referer": FORUM_URL,
        }
        if impersonate:
            from curl_cffi import requests as cffi          # noqa: PLC0415
            self._sess = cffi.Session(impersonate=impersonate, timeout=timeout)
        else:
            s = requests.Session()
            s.headers.update(self.headers)
            retry = Retry(
                total=4, connect=3, read=3,
                backoff_factor=1.5,
                status_forcelist=(429, 500, 502, 503, 504),
                allowed_methods=frozenset(["GET"]),
                respect_retry_after_header=True,
                raise_on_status=False,
            )
            s.mount("https://", HTTPAdapter(max_retries=retry, pool_maxsize=4))
            self._sess = s

    def _throttle(self) -> None:
        if not self.delay:
            return
        wait = self.delay + random.uniform(0, self.delay * 0.4)   # jitter
        elapsed = time.monotonic() - self._last
        if elapsed < wait:
            time.sleep(wait - elapsed)
        self._last = time.monotonic()

    def get_html(self, url: str) -> str:
        self._throttle()
        log.debug("GET %s", url)
        if self.impersonate:
            r = self._sess.get(url, headers=self.headers)
        else:
            r = self._sess.get(url, timeout=self.timeout)
        if r.status_code in (404, 410):
            raise ThreadGone(f"{r.status_code} — thread da bi xoa hoac chuyen: {url}")
        if r.status_code == 403:
            raise PermissionError(
                f"403 tu voz.vn cho {url}. IP hoac fingerprint bi Cloudflare chan. "
                "Thu: pip install curl_cffi && them --impersonate chrome, "
                "hoac chay tu may ca nhan / self-hosted runner."
            )
        r.raise_for_status()
        r.encoding = r.encoding or "utf-8"
        return r.text

    def close(self) -> None:
        try:
            self._sess.close()
        except Exception:
            pass


class ThreadGone(Exception):
    """Thread khong con truy cap duoc (404/410).

    Rat hay xay ra voi box Diem bao: mod xoa hoac chuyen thread trong khoang
    thoi gian giua luc doc danh sach page 1 va luc vao chi tiet. Day KHONG phai
    loi cua script — no chi khong con gi de quet.
    """


def robots_allows(url: str, ua: str) -> bool:
    """Kiem tra robots.txt cua voz.vn cho User-Agent dang dung."""
    rp = robotparser.RobotFileParser()
    rp.set_url(BASE + "/robots.txt")
    try:
        rp.read()
    except Exception as exc:
        log.warning("Khong doc duoc robots.txt (%s) -> coi nhu cho phep.", exc)
        return True
    return rp.can_fetch(ua, url)


# ------------------------------------------------------------------ parsing
def soup_of(html: str) -> BeautifulSoup:
    try:
        return BeautifulSoup(html, "lxml")
    except Exception:
        return BeautifulSoup(html, "html.parser")


def _abs(href: str) -> str:
    return urlparse.urljoin(BASE, href or "")


def normalize_thread_url(href: str) -> tuple[str, str]:
    """'/t/slug.123/unread' -> ('https://voz.vn/t/slug.123/', '123')"""
    path = urlparse.urlsplit(_abs(href)).path
    m = THREAD_URL_RE.match(path)
    if not m:
        return _abs(href), ""
    return f"{BASE}/t/{m.group('slug')}/", m.group("id")


def unwrap_proxy(url: str) -> str:
    """XenForo boc link ngoai qua /proxy.php?link=... -> lay link that."""
    parts = urlparse.urlsplit(url)
    if parts.path.endswith("/proxy.php"):
        q = urlparse.parse_qs(parts.query)
        for key in ("link", "url", "image"):
            if q.get(key):
                return urlparse.unquote(q[key][0])
    return url


def parse_thread_list(html: str) -> list[Thread]:
    """Lay thread thuong o page 1, BO sticky.

    voz (XenForo 2) chia thread thanh 2 khoi:
        div.structItemContainer-group--sticky   <- thread ghim (chu mau cam)
        div.structItemContainer-group.js-threadList <- thread thuong
    """
    soup = soup_of(html)
    group = soup.select_one("div.structItemContainer-group.js-threadList")
    if group is not None:
        items = group.select("div.structItem--thread")
    else:                                    # fallback neu voz doi markup
        log.warning("Khong thay .js-threadList, dung fallback loc sticky thu cong.")
        items = [
            it for it in soup.select("div.structItem--thread")
            if it.find_parent("div", class_="structItemContainer-group--sticky") is None
            and "is-sticky" not in (it.get("class") or [])
        ]

    threads: list[Thread] = []
    for it in items:
        title_cell = it.select_one(".structItem-title")
        if not title_cell:
            continue
        link = (title_cell.select_one('a[data-tp-primary="on"]')
                or next((a for a in title_cell.select("a[href]")
                         if "labelLink" not in (a.get("class") or [])), None))
        if not link:
            continue
        url, tid = normalize_thread_url(link.get("href", ""))
        if not tid:
            continue
        prefix = " ".join(
            a.get_text(strip=True) for a in title_cell.select("a.labelLink, .label")
        )
        threads.append(Thread(
            thread_id=tid,
            title=link.get_text(" ", strip=True),
            url=url,
            prefix=prefix,
        ))
    return threads


def last_page_number(soup: BeautifulSoup) -> int:
    nums = [1]
    for a in soup.select(".pageNav-page a, .pageNav-main a"):
        t = a.get_text(strip=True).replace(",", "")
        if t.isdigit():
            nums.append(int(t))
    jump = soup.select_one(".pageNavSimple-el--current")
    if jump:
        m = re.search(r"of\s+([\d,]+)", jump.get_text(" ", strip=True))
        if m:
            nums.append(int(m.group(1).replace(",", "")))
    return max(nums)


def parse_reactions(article) -> int:
    """Dem so luot 'Ung'.

    voz chi co 1 loai reaction (id=1, ten 'Ung'), nen day chinh la so luot ung.
    XenForo render: '<bdi>A</bdi>, <bdi>B</bdi>, <bdi>C</bdi> and 25 others'
      -> 3 ten hien + 25 = 28
    """
    link = (article.select_one(".message-footer .reactionsBar-link")
            or article.select_one(".reactionsBar-link"))
    if link is None:
        return 0
    named = len(link.select("bdi"))
    txt = link.get_text(" ", strip=True)
    if named == 0:
        named = 1 if txt else 0
    m = re.search(r"\band\s+([\d.,]+)\s+other", txt, re.I)
    others = int(re.sub(r"[.,\s]", "", m.group(1))) if m else 0
    return named + others


def post_text(article, keep_quotes: bool = False) -> str:
    """Lay noi dung post.

    keep_quotes=False (mac dinh, dung cho COMMENT): bo blockquote de chi giu loi
        cua chinh nguoi comment.
    keep_quotes=True (dung cho POST #1): giu quote, vi post #1 trong box Diem bao
        thuong CHI gom doan trich bai bao trong blockquote + card unfurl —
        strip quote se ra chuoi rong.
    """
    body = (article.select_one(".message-body .bbWrapper")
            or article.select_one(".message-body")
            or article.select_one(".message-content"))
    if body is None:
        return ""
    clone = BeautifulSoup(str(body), "html.parser")
    junk_sel = (
        "script, style, .bbCodeBlock--unfurl, .js-unfurl, .message-signature, "
        ".bbCodeSpoiler-button, .js-selectToQuoteEnd, .bbCodeBlock-expandLink, "
        ".bbMediaWrapper, noscript"
    )
    if not keep_quotes:
        junk_sel += ", blockquote, .bbCodeBlock--quote"
    for tag in clone.select(junk_sel):
        tag.decompose()
    for br in clone.select("br"):
        br.replace_with("\n")
    text = clone.get_text("\n", strip=True)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_article_url(article) -> str:
    """Tim link bai bao trong post #1.

    Uu tien card unfurl (XenForo tu sinh khi dan link bao), sau do link ngoai
    dau tien khong phai anh / khong phai voz.
    """
    body = article.select_one(".message-body") or article
    ordered = []
    for sel in ("a.bbCodeBlock-title[href]", ".bbCodeBlock--unfurl a[href]",
                ".js-unfurl a[href]"):
        ordered.extend(body.select(sel))
    ordered.extend(body.select("a[href]"))

    for a in ordered:
        raw = unwrap_proxy(_abs(a.get("href", "")))
        parts = urlparse.urlsplit(raw)
        if parts.scheme not in ("http", "https"):
            continue
        host = parts.netloc.lower()
        if host in NON_ARTICLE_HOSTS or host.endswith(".voz.vn"):
            continue
        if IMAGE_EXT_RE.search(parts.path):
            continue
        return raw
    return ""


def parse_post_time(article) -> Optional[datetime]:
    """Lay thoi diem dang post.

    voz/XenForo render:
        <time datetime="2026-09-03T09:02:03+0700" data-timestamp="1788400923"
              data-date="Sep 3, 2026" data-time="9:02 AM" data-short="2h">
    CHU Y: data-time la CHUOI GIO ("9:02 AM"), unix timestamp nam o data-timestamp.
    """
    tt = article.select_one(".message-attribution-main time") or article.select_one("time")
    if tt is None:
        return None
    for attr in ("data-timestamp", "data-time"):
        raw = (tt.get(attr) or "").strip()
        if raw.isdigit():
            try:
                return datetime.fromtimestamp(int(raw), tz=timezone.utc)
            except (ValueError, OSError, OverflowError):
                pass
    iso = (tt.get("datetime") or "").strip()
    if iso:
        iso = iso.replace("Z", "+00:00")
        # Python < 3.11 khong doc duoc offset dang "+0700" -> chuyen thanh "+07:00"
        iso = re.sub(r"([+-]\d{2})(\d{2})$", r"\1:\2", iso)
        try:
            return datetime.fromisoformat(iso)
        except ValueError:
            pass
    return None


def parse_posts(html: str, thread_url: str) -> tuple[list[Post], int, BeautifulSoup]:
    soup = soup_of(html)
    posts: list[Post] = []
    for art in soup.select("article.message--post"):
        content = art.get("data-content", "") or ""
        pid = content.rsplit("-", 1)[-1] if content else ""
        number = None
        for a in art.select(".message-attribution-opposite a"):
            t = a.get_text(strip=True)
            if t.startswith("#"):
                digits = re.sub(r"[^\d]", "", t)
                number = int(digits) if digits else None
        created = parse_post_time(art)
        posts.append(Post(
            post_id=pid,
            number=number,
            author=art.get("data-author", "") or "",
            url=f"{BASE}/p/{pid}/" if pid else thread_url,
            created_at=created,
            reactions=parse_reactions(art),
            text=post_text(art),
        ))
    return posts, last_page_number(soup), soup


# ------------------------------------------------------------------ crawling
def page_url(thread_url: str, page: int) -> str:
    return thread_url if page <= 1 else f"{thread_url.rstrip('/')}/page-{page}"


def scrape_thread(fetcher: Fetcher, th: Thread, *, max_pages: int = 0,
                  top_n: int = 20, min_reactions: int = 1,
                  since_hours: int = 0) -> Thread:
    """Quet HET cac trang cua 1 thread, tinh diem, lay top comment."""
    page = 1
    total_pages = 1
    seen: set[str] = set()
    all_posts: list[Post] = []

    while True:
        url = page_url(th.url, page)
        html = fetcher.get_html(url)
        posts, last, soup = parse_posts(html, th.url)
        if page == 1:
            total_pages = last
            first_art = soup.select_one("article.message--post")
            if first_art is not None:
                th.article_url = extract_article_url(first_art)
                # post #1 thuong chi gom quote bai bao -> giu quote de lay tom tat
                th.article_excerpt = post_text(first_art, keep_quotes=True)
                if th.article_url:
                    th.article_source = urlparse.urlsplit(th.article_url).netloc.lower()
                    th.article_source = re.sub(r"^www\.", "", th.article_source)
        for p in posts:
            if p.post_id and p.post_id in seen:
                continue
            if p.post_id:
                seen.add(p.post_id)
            all_posts.append(p)

        th.pages_scanned = page
        if page >= total_pages or (max_pages and page >= max_pages):
            break
        page += 1

    all_posts.sort(key=lambda p: (p.number is None, p.number or 0))
    th.posts = all_posts

    op = all_posts[0] if all_posts else None
    if op is not None:
        th.op_author = op.author
        th.op_created_at = op.created_at

    comments = [p for p in all_posts if (p.number or 0) != 1]
    th.replies = len(comments)
    th.total_reactions = sum(p.reactions for p in all_posts)

    pool = comments
    if since_hours:
        cutoff = datetime.now(timezone.utc) - timedelta(hours=since_hours)
        pool = [p for p in pool if p.created_at and p.created_at >= cutoff] or comments
    pool = [p for p in pool if p.reactions >= min_reactions and p.text]
    pool.sort(key=lambda p: (-p.reactions, p.created_at or datetime.max.replace(tzinfo=timezone.utc)))
    th.top_comments = pool[:top_n]

    th.base_score = W_REPLY * th.replies + W_UNG * th.total_reactions
    age_h = 0.0
    if th.op_created_at:
        age_h = max(0.0, (datetime.now(timezone.utc) - th.op_created_at).total_seconds() / 3600)
    th.hot_score = th.base_score / ((age_h + AGE_OFFSET_H) ** GRAVITY)
    return th


def crawl(fetcher: Fetcher, *, top_n: int, min_reactions: int, max_threads: int,
          max_pages: int, since_hours: int, dry_run: bool) -> list[Thread]:
    log.info("Doc danh sach thread: %s", FORUM_URL)
    threads = parse_thread_list(fetcher.get_html(FORUM_URL))
    log.info("Tim thay %d thread thuong o page 1 (da bo sticky).", len(threads))
    if not threads:
        raise RuntimeError(
            "Khong parse duoc thread nao. Co the voz doi markup, hoac bi chan. "
            "Chay lai voi --debug de xem chi tiet."
        )
    if max_threads:
        threads = threads[:max_threads]
    if dry_run:
        return threads

    done: list[Thread] = []
    for i, th in enumerate(threads, 1):
        log.info("[%d/%d] %s", i, len(threads), th.title[:70])
        try:
            scrape_thread(fetcher, th, max_pages=max_pages, top_n=top_n,
                          min_reactions=min_reactions, since_hours=since_hours)
            log.info("      -> %d trang, %d reply, %d Ung, hot=%.1f",
                     th.pages_scanned, th.replies, th.total_reactions, th.hot_score)
        except PermissionError:
            raise
        except ThreadGone as exc:
            th.status = "gone"
            th.error = str(exc)
            th.hot_score = 0.0
            log.warning("      -> BO QUA: %s", th.error)
        except Exception as exc:
            th.status = "error"
            th.error = f"{type(exc).__name__}: {exc}"
            th.hot_score = 0.0
            log.error("      -> LOI: %s", th.error)
        done.append(th)

    # Thread ok xep truoc theo hot score; gone/error luon xuong duoi cung
    # (van giu lai trong output de khong che mat van de).
    done.sort(key=lambda t: (t.status != "ok", -t.hot_score))
    return done


# ------------------------------------------------------------------ excel
FONT = "Arial"
HDR_FILL = "1F3864"
YELLOW = "FFFF00"
CFG_SHEET = "ThamSo"          # ten khong dau, khong khoang trang -> de tham chieu


def _style_header(ws, ncols: int, row: int = 1) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=HDR_FILL)
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _set_widths(ws, widths: dict[str, int]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _link(ws, row: int, col: int, url: str, label: str = "") -> None:
    from openpyxl.styles import Font
    cell = ws.cell(row=row, column=col, value=(label or url)[:255] if (label or url) else "")
    if url:
        cell.hyperlink = url
        cell.font = Font(name=FONT, size=10, color="0563C1", underline="single")
    else:
        cell.font = Font(name=FONT, size=10)


def _clip(text: str, n: int) -> str:
    text = (text or "").strip()
    return text if len(text) <= n else text[: n - 1] + "…"


def build_workbook(threads: list[Thread], top_n: int, run_at: datetime):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    body = Font(name=FONT, size=10)
    wb = Workbook()

    # ---------------- Sheet: ThamSo (tunable) ----------------
    cfg = wb.active
    cfg.title = CFG_SHEET
    cfg["A1"] = "Tham so tinh diem — SUA O CELL VANG, cac sheet khac tu tinh lai"
    cfg["A1"].font = Font(name=FONT, bold=True, size=11)
    rows = [
        ("Trong so moi reply (W_REPLY)", W_REPLY),
        ("Trong so moi luot Ung (W_UNG)", W_UNG),
        ("Do giam theo tuoi (GRAVITY, <1)", GRAVITY),
        ("Bu tuoi tinh theo gio (AGE_OFFSET_H)", AGE_OFFSET_H),
    ]
    for i, (label, val) in enumerate(rows, start=2):
        cfg.cell(row=i, column=1, value=label).font = body
        c = cfg.cell(row=i, column=2, value=val)
        c.font = Font(name=FONT, size=10, color="0000FF", bold=True)
        c.fill = PatternFill("solid", fgColor=YELLOW)
    cfg["A7"] = "Diem goc      = replies * W_REPLY + tong_Ung * W_UNG"
    cfg["A8"] = "Hot score     = Diem goc / (tuoi_thread_gio + AGE_OFFSET_H) ^ GRAVITY"
    cfg["A9"] = "Nguon du lieu = " + FORUM_URL + "  (page 1, da bo thread ghim)"
    cfg["A10"] = "Thoi diem quet= " + run_at.strftime("%Y-%m-%d %H:%M:%S %Z")
    cfg["A11"] = "Luu y: 'Ung' la reaction duy nhat cua voz (reaction id 1)."
    cfg["A12"] = ("Ban quyen: noi dung chi dung lam tu lieu tham khao/y tuong. "
                  "Khong copy nguyen van comment cua nguoi khac len Threads.")
    for r in range(7, 13):
        cfg.cell(row=r, column=1).font = Font(name=FONT, size=9, italic=True)
    _set_widths(cfg, {"A": 62, "B": 14})

    # ---------------- Sheet: Tong quan ----------------
    ov = wb.create_sheet("Tong quan")
    headers = ["Hang", "Tieu de", "Link thread", "Link bai bao", "Nguon",
               "Replies", "Tong Ung", "Tuoi thread (gio)", "Diem goc", "Hot score",
               "So comment top", "Nguoi dang", "Thoi gian dang", "Trang da quet",
               "Trang thai", "Trich bai bao (post #1)", "Loi"]
    ov.append(headers)
    _style_header(ov, len(headers))
    n = len(threads)
    for i, th in enumerate(threads):
        r = i + 2
        age_h = 0.0
        if th.op_created_at:
            age_h = max(0.0, (run_at.astimezone(timezone.utc) - th.op_created_at)
                        .total_seconds() / 3600)
        ov.cell(row=r, column=1, value=f"=RANK(J{r},$J$2:$J${n + 1})").font = body
        ov.cell(row=r, column=2, value=_clip(th.title, 300)).font = body
        _link(ov, r, 3, th.url, th.url)
        _link(ov, r, 4, th.article_url, th.article_url)
        ov.cell(row=r, column=5, value=th.article_source).font = body
        ov.cell(row=r, column=6, value=th.replies).font = body
        ov.cell(row=r, column=7, value=th.total_reactions).font = body
        ov.cell(row=r, column=8, value=round(age_h, 2)).font = body
        ov.cell(row=r, column=9,
                value=f"=F{r}*{CFG_SHEET}!$B$2+G{r}*{CFG_SHEET}!$B$3").font = body
        ov.cell(row=r, column=10,
                value=(f"=IFERROR(I{r}/POWER(H{r}+{CFG_SHEET}!$B$5,"
                       f"{CFG_SHEET}!$B$4),0)")).font = body
        ov.cell(row=r, column=11, value=len(th.top_comments)).font = body
        ov.cell(row=r, column=12, value=th.op_author).font = body
        ov.cell(row=r, column=13,
                value=(th.op_created_at.astimezone(TZ_VN).strftime("%Y-%m-%d %H:%M")
                       if th.op_created_at else "")).font = body
        ov.cell(row=r, column=14, value=th.pages_scanned).font = body
        st = ov.cell(row=r, column=15, value={
            "ok": "ok",
            "gone": "da xoa/chuyen (404)",
            "error": "loi",
        }.get(th.status, th.status))
        st.font = body
        if th.status != "ok":
            st.fill = PatternFill("solid", fgColor="FFE699")
        exc = ov.cell(row=r, column=16, value=_clip(th.article_excerpt, 4000))
        exc.font = body
        exc.alignment = Alignment(wrap_text=True, vertical="top")
        ov.cell(row=r, column=17, value=_clip(th.error, 300)).font = body
        for col in (8, 9, 10):
            ov.cell(row=r, column=col).number_format = "0.0"
        ov.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    _set_widths(ov, {"A": 6, "B": 52, "C": 34, "D": 40, "E": 16, "F": 9, "G": 10,
                     "H": 11, "I": 10, "J": 11, "K": 9, "L": 20, "M": 17, "N": 8,
                     "O": 19, "P": 70, "Q": 24})
    if n:
        ov.auto_filter.ref = f"A1:Q{n + 1}"

    # ---------------- Sheet: Top comments (dang dai — de doc/loc) ----------------
    dt = wb.create_sheet("Top comments")
    h2 = ["Hang thread", "Tieu de thread", "Link bai bao", "Hang comment", "Ung",
          "Tac gia", "Post #", "Link comment", "Thoi gian", "Noi dung comment"]
    dt.append(h2)
    _style_header(dt, len(h2))
    row = 2
    for ti, th in enumerate(threads, start=1):
        for ci, p in enumerate(th.top_comments, start=1):
            dt.cell(row=row, column=1, value=ti).font = body
            dt.cell(row=row, column=2, value=_clip(th.title, 300)).font = body
            _link(dt, row, 3, th.article_url, th.article_url)
            dt.cell(row=row, column=4, value=ci).font = body
            dt.cell(row=row, column=5, value=p.reactions).font = body
            dt.cell(row=row, column=6, value=p.author).font = body
            dt.cell(row=row, column=7, value=p.number).font = body
            _link(dt, row, 8, p.url, p.url)
            dt.cell(row=row, column=9,
                    value=(p.created_at.astimezone(TZ_VN).strftime("%Y-%m-%d %H:%M")
                           if p.created_at else "")).font = body
            c = dt.cell(row=row, column=10, value=_clip(p.text, 4000))
            c.font = body
            c.alignment = Alignment(wrap_text=True, vertical="top")
            dt.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
    _set_widths(dt, {"A": 11, "B": 40, "C": 34, "D": 12, "E": 7, "F": 20, "G": 8,
                     "H": 30, "I": 17, "J": 95})
    if row > 2:
        dt.auto_filter.ref = f"A1:J{row - 1}"

    # ---------------- Sheet: Bang ngang (1 thread = 1 dong, Top 1..N theo cot) ----
    wide = wb.create_sheet("Bang ngang")
    h3 = (["Hang", "Tieu de", "Link bai bao", "Link thread", "Replies", "Tong Ung"]
          + [f"Top {i}" for i in range(1, top_n + 1)])
    wide.append(h3)
    _style_header(wide, len(h3))
    for i, th in enumerate(threads, start=1):
        r = i + 1
        wide.cell(row=r, column=1, value=i).font = body
        wide.cell(row=r, column=2, value=_clip(th.title, 300)).font = body
        _link(wide, r, 3, th.article_url, th.article_url)
        _link(wide, r, 4, th.url, th.url)
        wide.cell(row=r, column=5, value=th.replies).font = body
        wide.cell(row=r, column=6, value=th.total_reactions).font = body
        for j in range(top_n):
            cell = wide.cell(row=r, column=7 + j)
            if j < len(th.top_comments):
                p = th.top_comments[j]
                cell.value = _clip(f"[{p.reactions} Ung] @{p.author}: {p.text}", 900)
            cell.font = body
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        wide.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        wide.row_dimensions[r].height = 90
    widths = {"A": 6, "B": 40, "C": 30, "D": 30, "E": 9, "F": 10}
    from openpyxl.utils import get_column_letter
    for j in range(top_n):
        widths[get_column_letter(7 + j)] = 46
    _set_widths(wide, widths)

    return wb


# ------------------------------------------------------------------ main
def parse_args(argv: Optional[list[str]] = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Quet box Diem bao voz.vn -> xep hang thread hot + top comment -> Excel",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    p.add_argument("--out-dir", default="output", help="Thu muc xuat file")
    p.add_argument("--top", type=int, default=20, help="So comment top moi thread")
    p.add_argument("--min-reactions", type=int, default=1,
                   help="Chi lay comment co >= bao nhieu luot Ung")
    p.add_argument("--delay", type=float, default=1.2,
                   help="Giay nghi toi thieu giua 2 request (lich su voi server)")
    p.add_argument("--max-threads", type=int, default=0,
                   help="Gioi han so thread (0 = tat ca thread thuong o page 1)")
    p.add_argument("--max-pages", type=int, default=0,
                   help="Gioi han so trang moi thread (0 = quet het)")
    p.add_argument("--since-hours", type=int, default=0,
                   help="Chi xet comment dang trong N gio qua (0 = tat ca)")
    p.add_argument("--user-agent", default=DEFAULT_UA)
    p.add_argument("--impersonate", default="",
                   help="Dung curl_cffi gia lap TLS, vd: chrome (khi bi 403)")
    p.add_argument("--ignore-robots", action="store_true",
                   help="Bo qua kiem tra robots.txt (khong khuyen khich)")
    p.add_argument("--dry-run", action="store_true",
                   help="Chi liet ke thread page 1, khong vao chi tiet")
    p.add_argument("--no-json", action="store_true", help="Khong xuat file JSON")
    p.add_argument("--debug", action="store_true")
    return p.parse_args(argv)


def main(argv: Optional[list[str]] = None) -> int:
    args = parse_args(argv)
    logging.basicConfig(
        level=logging.DEBUG if args.debug else logging.INFO,
        format="%(asctime)s %(levelname)-7s %(message)s",
        datefmt="%H:%M:%S",
        stream=sys.stderr,
    )

    if not args.ignore_robots:
        if not robots_allows(FORUM_URL, args.user_agent):
            log.error("robots.txt cua voz.vn KHONG cho phep User-Agent nay doc %s.\n"
                      "Doi User-Agent (--user-agent) hoac dung --ignore-robots neu ban "
                      "chac chan minh duoc phep.", FORUM_URL)
            return 2
        log.info("robots.txt: OK (duoc phep doc %s).", FORUM_PATH)

    run_at = datetime.now(TZ_VN)
    out_dir = os.path.abspath(args.out_dir)
    os.makedirs(out_dir, exist_ok=True)

    fetcher = Fetcher(ua=args.user_agent, delay=args.delay, impersonate=args.impersonate)
    try:
        threads = crawl(
            fetcher,
            top_n=args.top,
            min_reactions=args.min_reactions,
            max_threads=args.max_threads,
            max_pages=args.max_pages,
            since_hours=args.since_hours,
            dry_run=args.dry_run,
        )
    except PermissionError as exc:
        log.error("%s", exc)
        return 3
    finally:
        fetcher.close()

    if args.dry_run:
        for i, th in enumerate(threads, 1):
            print(f"{i:2d}. {th.title}\n    {th.url}")
        return 0

    stamp = run_at.strftime("%Y-%m-%d")
    xlsx_path = os.path.join(out_dir, f"voz-diem-bao_{stamp}.xlsx")
    wb = build_workbook(threads, args.top, run_at)
    wb.save(xlsx_path)
    log.info("Da ghi Excel: %s", xlsx_path)

    if not args.no_json:
        json_path = os.path.join(out_dir, f"voz-diem-bao_{stamp}.json")
        payload = {
            "source": FORUM_URL,
            "scraped_at": run_at.isoformat(),
            "weights": {"W_REPLY": W_REPLY, "W_UNG": W_UNG,
                        "GRAVITY": GRAVITY, "AGE_OFFSET_H": AGE_OFFSET_H},
            "threads": [t.to_dict() for t in threads],
        }
        with open(json_path, "w", encoding="utf-8") as fh:
            json.dump(payload, fh, ensure_ascii=False, indent=1)
        log.info("Da ghi JSON : %s", json_path)

    print("\n=== TOP 5 CHU DE HOT NHAT ===")
    for i, th in enumerate(threads[:5], 1):
        print(f"{i}. [hot {th.hot_score:6.1f} | {th.replies:4d} reply | "
              f"{th.total_reactions:5d} Ung] {th.title[:70]}")
        if th.top_comments:
            top = th.top_comments[0]
            print(f"   comment top: [{top.reactions} Ung] @{top.author}: "
                  f"{_clip(top.text, 120)}")
    gone = [t for t in threads if t.status == "gone"]
    failed = [t for t in threads if t.status == "error"]
    if gone:
        log.warning("%d thread da bi xoa/chuyen trong luc quet (binh thuong voi "
                    "box Diem bao): %s", len(gone),
                    "; ".join(t.title[:40] for t in gone))
    if failed:
        log.warning("%d thread bi LOI THAT (xem cot 'Loi' trong Excel): %s",
                    len(failed), "; ".join(t.title[:40] for t in failed))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
