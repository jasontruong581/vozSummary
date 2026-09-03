#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Test cho scrape_voz.py — fixture copy tu markup that cua voz.vn (XenForo 2).

Chay:  python -m pytest test_parser.py -q      (hoac: python test_parser.py)
"""
from datetime import datetime, timezone

from scrape_voz import (
    W_REPLY, W_UNG, GRAVITY, AGE_OFFSET_H,
    Post, Thread, extract_article_url, last_page_number, normalize_thread_url,
    ThreadGone, crawl, parse_post_time, parse_posts, parse_reactions,
    parse_thread_list, post_text, soup_of, unwrap_proxy,
)

# --------------------------------------------------------------------- fixtures
FORUM_HTML = """
<div class="structItemContainer">
 <div class="structItemContainer-group structItemContainer-group--sticky">
  <div class="structItem structItem--thread js-threadListItem-1216621">
   <div class="structItem-cell structItem-cell--main"><div class="structItem-title">
     <a href="/t/tinh-hinh-iran.1216621/unread" data-tp-primary="on">Tinh hinh Iran</a>
   </div></div>
   <div class="structItem-cell structItem-cell--meta">
     <dl class="pairs"><dt>Replies</dt><dd>22K</dd></dl></div>
  </div>
  <div class="structItem structItem--thread js-threadListItem-617079">
   <div class="structItem-cell structItem-cell--main"><div class="structItem-title">
     <a href="/t/report-f33.617079/unread" data-tp-primary="on">Report F33</a>
   </div></div>
  </div>
 </div>
 <div class="structItemContainer-group js-threadList">
  <div class="structItem structItem--thread is-unread js-threadListItem-1275581">
   <div class="structItem-cell structItem-cell--main"><div class="structItem-title">
     <a href="/f/kinh-te.35/" class="labelLink"><span class="label">Kinh te</span></a>
     <a href="/t/nguoi-tre-co-dang-tu-bo-viec-so-huu-nha.1275581/unread"
        data-tp-primary="on">Nguoi tre co dang tu bo viec so huu nha?</a>
   </div></div>
   <div class="structItem-cell structItem-cell--meta">
     <dl class="pairs"><dt>Replies</dt><dd>68</dd></dl></div>
  </div>
  <div class="structItem structItem--thread js-threadListItem-1275609">
   <div class="structItem-cell structItem-cell--main"><div class="structItem-title">
     <a href="/t/internet-viet-nam-dat-ky-luc-moi.1275609/" data-tp-primary="on">Internet Viet Nam dat ky luc moi</a>
   </div></div>
  </div>
 </div>
</div>
"""

THREAD_HTML = """
<nav class="pageNavWrapper"><div class="pageNav"><ul class="pageNav-main">
  <li class="pageNav-page pageNav-page--current"><a href="/t/x.1/">1</a></li>
  <li class="pageNav-page"><a href="/t/x.1/page-2">2</a></li>
  <li class="pageNav-page"><a href="/t/x.1/page-3">3</a></li>
  <li class="pageNav-page"><a href="/t/x.1/page-4">4</a></li>
</ul></div></nav>

<article class="message message--post js-post" data-author="MotVoHaiCon"
         data-content="post-43555707" id="js-post-43555707">
  <div class="message-attribution">
    <ul class="message-attribution-main"><li><a href="/t/x.1/post-43555707"
      ><time class="u-dt" dir="auto" datetime="2026-09-03T09:02:03+0700"
        data-timestamp="1788400923" data-date="Sep 3, 2026" data-time="9:02 AM"
        data-short="2h" title="Sep 3, 2026 at 9:02 AM">Today at 9:02 AM</time></a></li></ul>
    <ul class="message-attribution-opposite"><li><a href="/goto/post?id=43555707">#1</a></li></ul>
  </div>
  <div class="message-content">
   <div class="message-body"><div class="bbWrapper">
    <blockquote class="bbCodeBlock bbCodeBlock--quote"><div class="bbCodeBlock-content">
      (VTC News) - Gia nha tang cao khien nhieu nguoi tre kho khan.
      <a class="bbCodeBlock-expandLink" href="#">Click to expand...</a>
    </div></blockquote>
    Ae thay sao? Minh thi thue nha cho thoai mai.<br>Mua nha o HN gio kho qua.
    <div class="bbCodeBlock bbCodeBlock--unfurl js-unfurl">
      <a href="https://vtcnews.vn/nguoi-tre-co-dang-tu-bo-viec-so-huu-nha-ar1036654.html"
         class="bbCodeBlock-title">Nguoi tre co dang tu bo viec so huu nha?</a>
      <span class="bbCodeBlock-unfurlUrl">vtcnews.vn</span>
    </div>
   </div></div>
  </div>
  <footer class="message-footer"><div class="reactionsBar js-reactionsList is-active">
    <ul class="reactionSummary"><li><span class="reaction reaction--1" data-reaction-id="1"
      ><img class="reaction-image" alt="Ung" title="Ung"></span></li></ul>
    <a class="reactionsBar-link" href="/p/43555707/reactions"><bdi>itisme</bdi> and <bdi>Meow</bdi></a>
  </div></footer>
</article>

<article class="message message--post js-post" data-author="Holyshjt"
         data-content="post-43555800">
  <div class="message-attribution">
    <ul class="message-attribution-main"><li><time datetime="2026-09-03T09:07:00+0700"
      data-time="9:07 AM">Today at 9:07 AM</time></li></ul>
    <ul class="message-attribution-opposite"><li><a href="/goto/post?id=43555800">#2</a></li></ul>
  </div>
  <div class="message-content"><div class="message-body"><div class="bbWrapper">
    Gia nha gap 40 lan thu nhap thi ai mua duoc.
  </div></div></div>
  <footer class="message-footer"><div class="reactionsBar js-reactionsList is-active">
    <a class="reactionsBar-link" href="/p/43555800/reactions"
      ><bdi>A</bdi>, <bdi>B</bdi>, <bdi>C</bdi> and 1,208 others</a>
  </div></footer>
</article>

<article class="message message--post js-post" data-author="NoReact"
         data-content="post-43555801">
  <div class="message-attribution">
    <ul class="message-attribution-opposite"><li><a href="#">#3</a></li></ul>
  </div>
  <div class="message-content"><div class="message-body"><div class="bbWrapper">
    up cho bac
  </div></div></div>
  <footer class="message-footer"></footer>
</article>

<article class="message message--post js-post" data-author="OneReact"
         data-content="post-43555802">
  <div class="message-attribution">
    <ul class="message-attribution-opposite"><li><a href="#">#4</a></li></ul>
  </div>
  <div class="message-content"><div class="message-body"><div class="bbWrapper">
    Thue nha la lua chon hop ly voi nguoi tre.
  </div></div></div>
  <footer class="message-footer"><div class="reactionsBar is-active">
    <a class="reactionsBar-link" href="/p/43555802/reactions"><bdi>solo</bdi></a>
  </div></footer>
</article>
"""


# --------------------------------------------------------------------- tests
def test_thread_list_skips_sticky():
    threads = parse_thread_list(FORUM_HTML)
    assert [t.thread_id for t in threads] == ["1275581", "1275609"], \
        "phai bo 2 thread sticky, chi lay 2 thread thuong"
    assert threads[0].title == "Nguoi tre co dang tu bo viec so huu nha?"
    assert threads[0].url == "https://voz.vn/t/nguoi-tre-co-dang-tu-bo-viec-so-huu-nha.1275581/"
    assert "Kinh te" in threads[0].prefix          # prefix tach rieng khoi tieu de


def test_thread_list_fallback_without_group_class():
    """Neu voz doi markup, fallback van phai loc duoc sticky."""
    html = FORUM_HTML.replace('structItemContainer-group js-threadList',
                              'structItemContainer-group')
    threads = parse_thread_list(html)
    assert [t.thread_id for t in threads] == ["1275581", "1275609"]


def test_normalize_thread_url():
    assert normalize_thread_url("/t/abc.123/unread") == ("https://voz.vn/t/abc.123/", "123")
    assert normalize_thread_url("/t/abc.123/post-999") == ("https://voz.vn/t/abc.123/", "123")
    assert normalize_thread_url("/t/abc.123/page-7") == ("https://voz.vn/t/abc.123/", "123")


def test_reaction_counting():
    soup = soup_of(THREAD_HTML)
    arts = soup.select("article.message--post")
    assert parse_reactions(arts[0]) == 2            # 2 <bdi>, khong co "others"
    assert parse_reactions(arts[1]) == 1211         # 3 ten + 1,208 others
    assert parse_reactions(arts[2]) == 0            # khong co reactionsBar
    assert parse_reactions(arts[3]) == 1            # 1 nguoi duy nhat


def test_last_page_number():
    assert last_page_number(soup_of(THREAD_HTML)) == 4
    assert last_page_number(soup_of("<div></div>")) == 1


def test_post_text_strips_quote_and_unfurl():
    arts = soup_of(THREAD_HTML).select("article.message--post")
    txt = post_text(arts[0])
    assert "Ae thay sao?" in txt and "Mua nha o HN" in txt
    assert "VTC News" not in txt, "phai bo quote"
    assert "vtcnews.vn" not in txt, "phai bo card unfurl"
    assert "Click to expand" not in txt


def test_extract_article_url():
    arts = soup_of(THREAD_HTML).select("article.message--post")
    assert extract_article_url(arts[0]) == \
        "https://vtcnews.vn/nguoi-tre-co-dang-tu-bo-viec-so-huu-nha-ar1036654.html"


def test_extract_article_url_ignores_voz_and_images():
    html = """<article class="message--post"><div class="message-body"><div class="bbWrapper">
      <a href="/t/khac.999/">thread khac</a>
      <a href="https://i.imgur.com/a.png">anh</a>
      <a href="https://tuoitre.vn/bai-viet-2026.htm">bai bao</a>
    </div></div></article>"""
    art = soup_of(html).select_one("article")
    assert extract_article_url(art) == "https://tuoitre.vn/bai-viet-2026.htm"


def test_extract_article_url_empty_when_none():
    html = """<article class="message--post"><div class="message-body"><div class="bbWrapper">
      chi co chu, khong co link</div></div></article>"""
    assert extract_article_url(soup_of(html).select_one("article")) == ""


def test_unwrap_proxy():
    assert unwrap_proxy(
        "https://voz.vn/proxy.php?link=https%3A%2F%2Fvnexpress.net%2Fa.html&hash=x"
    ) == "https://vnexpress.net/a.html"
    assert unwrap_proxy("https://vnexpress.net/a.html") == "https://vnexpress.net/a.html"


def test_parse_posts_fields():
    posts, last, _ = parse_posts(THREAD_HTML, "https://voz.vn/t/x.1/")
    assert last == 4
    assert [p.number for p in posts] == [1, 2, 3, 4]
    assert posts[0].author == "MotVoHaiCon"
    assert posts[0].post_id == "43555707"
    assert posts[0].url == "https://voz.vn/p/43555707/"
    assert posts[0].created_at == datetime.fromtimestamp(1788400923, tz=timezone.utc)
    assert posts[1].reactions == 1211


def test_top_comment_selection_and_score():
    """Kiem tra logic loc top comment + cong thuc diem (khong can mang)."""
    posts, _, _ = parse_posts(THREAD_HTML, "https://voz.vn/t/x.1/")
    comments = [p for p in posts if p.number != 1]
    pool = sorted((p for p in comments if p.reactions >= 1),
                  key=lambda p: -p.reactions)
    assert [p.author for p in pool] == ["Holyshjt", "OneReact"], \
        "comment 0 Ung bi loai, con lai xep giam dan theo Ung"

    replies, total_ung, age_h = 68, 1500, 10.0
    base = W_REPLY * replies + W_UNG * total_ung
    hot = base / ((age_h + AGE_OFFSET_H) ** GRAVITY)
    assert base == 3068.0
    assert 300 < hot < 400            # sanity: thread hot -> diem lon nhung co giam theo tuoi

    # thread cu hon phai co diem thap hon khi cung base
    hot_old = base / ((240.0 + AGE_OFFSET_H) ** GRAVITY)
    assert hot_old < hot


def test_workbook_builds_and_has_expected_sheets(tmp_path=None):
    import tempfile, os
    from openpyxl import load_workbook
    from scrape_voz import build_workbook, TZ_VN

    posts, _, _ = parse_posts(THREAD_HTML, "https://voz.vn/t/x.1/")
    th = Thread(
        thread_id="1275581", title="Nguoi tre co dang tu bo viec so huu nha?",
        url="https://voz.vn/t/nguoi-tre.1275581/",
        article_url="https://vtcnews.vn/a.html", article_source="vtcnews.vn",
        op_author="MotVoHaiCon", op_created_at=posts[0].created_at,
        pages_scanned=4, replies=3, total_reactions=1214,
        posts=posts, top_comments=[p for p in posts[1:] if p.reactions >= 1],
    )
    th.base_score = W_REPLY * th.replies + W_UNG * th.total_reactions
    th.hot_score = th.base_score / (AGE_OFFSET_H ** GRAVITY)

    wb = build_workbook([th], top_n=20, run_at=datetime.now(TZ_VN))
    out = os.path.join(tempfile.mkdtemp(), "t.xlsx")
    wb.save(out)

    got = load_workbook(out)
    assert got.sheetnames == ["ThamSo", "Tong quan", "Top comments", "Bang ngang"]
    ov = got["Tong quan"]
    assert ov["B2"].value == "Nguoi tre co dang tu bo viec so huu nha?"
    assert ov["I2"].value.startswith("=F2*ThamSo!$B$2")     # diem goc la CONG THUC
    assert ov["J2"].value.startswith("=IFERROR(")           # hot score la CONG THUC
    assert ov["A2"].value == "=RANK(J2,$J$2:$J$2)"
    assert got["Bang ngang"]["G2"].value.startswith("[1211 Ung] @Holyshjt:")
    assert got["Top comments"]["E2"].value == 1211


# ---------------------------- regression: 2 bug phat hien tren HTML that -----
def test_data_time_is_a_clock_string_not_a_timestamp():
    """REGRESSION: voz render data-time="9:02 AM" (chuoi gio), KHONG phai unix.

    Unix timestamp nam o data-timestamp. Neu doc data-time se ra created_at=None
    -> tuoi thread = 0 -> hot score sai hoan toan.
    """
    arts = soup_of(THREAD_HTML).select("article.message--post")
    got = parse_post_time(arts[0])
    assert got == datetime.fromtimestamp(1788400923, tz=timezone.utc)


def test_post_time_falls_back_to_datetime_attribute():
    """Post khong co data-timestamp -> phai doc attr datetime (offset '+0700')."""
    arts = soup_of(THREAD_HTML).select("article.message--post")
    got = parse_post_time(arts[1])
    assert got is not None and got.hour == 9 and got.minute == 7
    assert got.utcoffset().total_seconds() == 7 * 3600


def test_post_time_none_when_no_time_tag():
    art = soup_of('<article class="message--post"></article>').select_one("article")
    assert parse_post_time(art) is None


def test_op_needs_keep_quotes_to_have_any_content():
    """REGRESSION: post #1 box Diem bao thuong CHI gom quote bai bao + card unfurl.

    Strip quote nhu comment -> chuoi rong -> mat het noi dung bai bao.
    """
    arts = soup_of(THREAD_HTML).select("article.message--post")
    op = arts[0]
    with_quote = post_text(op, keep_quotes=True)
    assert "Gia nha tang cao" in with_quote, "phai giu duoc doan trich bai bao"
    assert "vtcnews.vn" not in with_quote, "van phai bo card unfurl"
    # comment thi nguoc lai: bo quote de chi con loi cua chinh nguoi comment
    assert "Gia nha tang cao" not in post_text(op, keep_quotes=False)


def test_op_only_quote_gives_empty_comment_text():
    """Post chi co quote -> text comment rong (dung), nhung excerpt van co."""
    html = ('<article class="message--post"><div class="message-body"><div class="bbWrapper">'
            '<blockquote class="bbCodeBlock--quote"><div>noi dung bai bao</div></blockquote>'
            '</div></div></article>')
    art = soup_of(html).select_one("article")
    assert post_text(art) == ""
    assert "noi dung bai bao" in post_text(art, keep_quotes=True)


# ------------------- thread bi xoa giua luc quet (gap that o run #1) --------
class _StubFetcher:
    """Fetcher gia: thread 1275581 con song, 1275609 da bi xoa (404)."""

    def __init__(self):
        self.calls = []

    def get_html(self, url):
        self.calls.append(url)
        if url.endswith("/f/diem-bao.33/"):
            return FORUM_HTML
        if "1275609" in url:
            raise ThreadGone("404 — thread da bi xoa hoac chuyen: " + url)
        return THREAD_HTML

    def close(self):
        pass


def _crawl_stub():
    return crawl(_StubFetcher(), top_n=20, min_reactions=1, max_threads=0,
                 max_pages=0, since_hours=0, dry_run=False)


def test_deleted_thread_marked_gone_not_error():
    """REGRESSION (run #1 tren GitHub Actions): mod xoa thread trong luc quet.

    404 phai thanh status 'gone' — khong phai 'error' — vi day khong phai loi
    cua script. Truoc day no vao cot Loi voi hot=0 va nam lan giua bang xep hang.
    """
    threads = _crawl_stub()
    by_id = {t.thread_id: t for t in threads}
    assert by_id["1275609"].status == "gone"
    assert by_id["1275609"].hot_score == 0.0
    assert "404" in by_id["1275609"].error
    assert by_id["1275581"].status == "ok"
    assert by_id["1275581"].hot_score > 0


def test_gone_thread_sorts_last():
    threads = _crawl_stub()
    assert threads[0].status == "ok"
    assert threads[-1].status == "gone", "thread da xoa phai xuong duoi cung"


def test_gone_thread_does_not_abort_the_run():
    """Mot thread 404 khong duoc lam chet ca lan chay."""
    threads = _crawl_stub()
    assert len(threads) == 2
    assert any(t.status == "ok" for t in threads)


def test_overview_sheet_has_status_column():
    import os, tempfile
    from openpyxl import load_workbook
    from scrape_voz import build_workbook, TZ_VN

    threads = _crawl_stub()
    wb = build_workbook(threads, top_n=20, run_at=datetime.now(TZ_VN))
    out = os.path.join(tempfile.mkdtemp(), "s.xlsx")
    wb.save(out)
    ov = load_workbook(out)["Tong quan"]
    headers = [ov.cell(row=1, column=c).value for c in range(1, 18)]
    assert "Trang thai" in headers
    col = headers.index("Trang thai") + 1
    states = {ov.cell(row=r, column=col).value for r in (2, 3)}
    assert states == {"ok", "da xoa/chuyen (404)"}


if __name__ == "__main__":
    import sys, traceback
    fns = [(k, v) for k, v in sorted(globals().items()) if k.startswith("test_")]
    bad = 0
    for name, fn in fns:
        try:
            fn()
            print(f"PASS  {name}")
        except Exception:
            bad += 1
            print(f"FAIL  {name}")
            traceback.print_exc()
    print(f"\n{len(fns) - bad}/{len(fns)} passed")
    sys.exit(1 if bad else 0)
